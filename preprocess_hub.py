"""Clean the Global AMR R&D Hub export and total investment by country."""
# Copyright 2026 Tavpritesh Sethi and Jasmine Kaur (Tavlab, IIITD)
# Licensed under the Apache License, Version 2.0 (see LICENSE)

import collections
import csv
import datetime as dt
import hashlib
import json
import openpyxl
import os
from common import ATLAS_ALIASES, DROP_BY_RULE, NAME_TO_ISO3, ROOT


def clean_projects():
    """Parse project amounts, dates and funder/host countries."""
    # preprocess the Global AMR R&D Hub export

    HUB_ALIASES = {
        "Korea, Republic of": "KOR",
        "Moldova, Republic of": "MDA",
        "Tanzania, United Republic of": "TZA",
        "St. Vincent and the Grenadines": "VCT",
        "Bolivia": "BOL",
    }

    # Entries that are not sovereign states: (status, explanation)
    NON_STATE_ENTITIES = {
        "European Union": (
            "supranational",
            "EU-level funding body (Horizon 2020/Europe, JPIAMR, IMI); cannot be attributed "
            "to a single country without an allocation assumption.",
        ),
        "Global Partnership": (
            "supranational",
            "Multilateral/global funding partnership; not a country.",
        ),
        "None": (
            "unattributed",
            "Literal string 'None' in the export: country not recorded.",
        ),
        "Denmark, South Korea": (
            "multi_country",
            "Two funder countries in one cell; the export gives no rule for splitting the amount.",
        ),
        "Kosovo": (
            "disputed_no_iso3",
            "No officially assigned ISO 3166-1 alpha-3 code; not a WHO Member State.",
        ),
        "Palestinian Territory, Occupied": (
            "non_member_state",
            "ISO3 PSE retained for completeness; does not join to TrACSS.",
        ),
        "Reunion": (
            "subnational_territory",
            "French overseas department (ISO3 REU); retained as REU, not folded into FRA.",
        ),
    }

    NON_STATE_ISO3 = {
        "Palestinian Territory, Occupied": "PSE",
        "Reunion": "REU",
    }

    OUT = "outputs/hub"
    os.makedirs("outputs/hub", exist_ok=True)
    SOURCE_XLSX = "data/hub/R&D.xlsx"
    RETRIEVAL_DATE = "2026-07-20"   # from the export's metadata sheet

    SOURCE_URLS = {
        "hub_dashboard_home": "https://dashboard.globalamrhub.org/",
        "hub_project_export_page": "https://dashboard.globalamrhub.org/public/projectdata",
        "hub_definitions": "https://globalamrhub.org/dynamic-dashboard/library/categories-and-definitions/",
    }

    def resolve_country(raw):
        if raw is None:
            return None, "", "blank"
        name = " ".join(str(raw).split())
        if name == "" or name.lower() == "nan":
            return None, name, "blank"
        if name in DROP_BY_RULE:
            return None, name, "drop_by_rule"
        if name in NON_STATE_ENTITIES:
            return NON_STATE_ISO3.get(name), name, NON_STATE_ENTITIES[name][0]
        for table in (NAME_TO_ISO3, HUB_ALIASES, ATLAS_ALIASES):
            if name in table:
                return table[name], name, "mapped"
        return None, name, "UNMAPPED"

    def parse_amount(raw):
        if raw is None:
            return None, "blank", ""
        if isinstance(raw, bool):
            return None, "unparsed", repr(raw)
        if isinstance(raw, (int, float)):
            return float(raw), "ok", repr(raw)
        s = str(raw).strip()
        if s == "":
            return None, "blank", s
        cleaned = s.replace(",", "").replace("$", "").replace("€", "").strip()
        try:
            return float(cleaned), "ok", s
        except ValueError:
            return None, "unparsed", s

    def parse_year(raw):
        if raw is None:
            return None, "blank", ""
        if isinstance(raw, int):
            return raw, "ok", str(raw)
        s = str(raw).strip()
        if s == "":
            return None, "blank", s
        try:
            return int(float(s)), "ok", s
        except ValueError:
            return None, "unparsed", s

    def sha256(path):
        h = hashlib.sha256()
        with open(path, "rb") as f:
            for chunk in iter(lambda: f.read(1 << 20), b""):
                h.update(chunk)
        return h.hexdigest()

    run_started = dt.datetime.now().isoformat(timespec="seconds")

    wb = openpyxl.load_workbook(SOURCE_XLSX, read_only=True, data_only=True)
    meta = {}
    for r in wb["metadata"].iter_rows(values_only=True):
        if r and r[0]:
            meta[str(r[0])] = r[1]

    ws = wb["data"]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    HEADER_ROW = 1
    col = {h: i for i, h in enumerate(header)}
    rows = list(it)
    n = len(rows)

    # data dictionary
    dict_rows = []
    for i, h in enumerate(header):
        vals = [r[i] for r in rows]
        nonnull = [v for v in vals if v is not None and str(v).strip() != ""]
        uniq = collections.Counter(str(v) for v in nonnull)
        types = collections.Counter(type(v).__name__ for v in vals)
        top = "; ".join(f"{k} (n={c})" for k, c in uniq.most_common(8) if len(k) < 60)
        dict_rows.append({
            "column_index": i, "column_name": h,
            "python_types": "; ".join(f"{k}={c}" for k, c in types.most_common()),
            "n_rows": n, "n_nonnull": len(nonnull),
            "pct_complete": round(100.0 * len(nonnull) / n, 2),
            "n_unique": len(uniq), "top_values": top,
        })
    with open(os.path.join(OUT, "hub_data_dictionary.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(dict_rows[0].keys()))
        w.writeheader()
        w.writerows(dict_rows)

    # country audit
    audit, unmapped = {}, []
    for field in ("Funder Country", "Institution Country"):
        counts = collections.Counter(r[col[field]] for r in rows)
        for raw, cnt in counts.items():
            iso3, canon, status = resolve_country(raw)
            audit[(field, canon)] = {
                "field": field, "raw_value": "" if raw is None else str(raw),
                "canonical_name": canon, "iso3": iso3 or "", "status": status,
                "n_projects": cnt, "rule_note": NON_STATE_ENTITIES.get(canon, ("", ""))[1],
            }
            if status == "UNMAPPED":
                unmapped.append((field, canon, cnt))
    with open(os.path.join(OUT, "hub_country_mapping_audit.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["field", "raw_value", "canonical_name", "iso3",
                                          "status", "n_projects", "rule_note"])
        w.writeheader()
        for k in sorted(audit):
            w.writerow(audit[k])
    if unmapped:
        raise AssertionError("Unmapped country names:\n" + "\n".join(
            f"  {f}: {nm!r} (n={c})" for f, nm, c in unmapped))

    # harmonised project table
    unparsed_log = []
    proj_fields = [
        "hub_project_id", "source_row_index",
        "funder_name", "funder_country_raw", "funder_iso3", "funder_country_status",
        "institution_name", "institution_country_raw", "institution_iso3",
        "institution_country_status",
        "start_year", "start_year_status", "end_year", "end_year_status", "duration_years",
        "amount_usd", "amount_usd_status", "amount_usd_raw",
        "amount_eur", "amount_eur_status", "amount_eur_raw",
        "sector", "sector_subcategory", "research_area", "research_subcategory",
        "disease", "infectious_agent", "individual_infectious_agent", "product_name",
        "project_group", "acronym",
        "source_file", "source_sheet", "header_row", "retrieval_date",
    ]
    out_rows = []
    for ridx, r in enumerate(rows):
        excel_row = ridx + 1 + HEADER_ROW
        f_iso3, _, f_status = resolve_country(r[col["Funder Country"]])
        i_iso3, _, i_status = resolve_country(r[col["Institution Country"]])
        usd, usd_st, usd_raw = parse_amount(r[col["Amount USD"]])
        eur, eur_st, eur_raw = parse_amount(r[col["Amount EUR"]])
        sy, sy_st, _ = parse_year(r[col["Start Year"]])
        ey, ey_st, _ = parse_year(r[col["End Year"]])

        for cname, st, rawv in (("Amount USD", usd_st, usd_raw), ("Amount EUR", eur_st, eur_raw),
                                ("Start Year", sy_st, str(r[col["Start Year"]])),
                                ("End Year", ey_st, str(r[col["End Year"]]))):
            if st == "unparsed":
                unparsed_log.append({
                    "source_file": os.path.basename(SOURCE_XLSX), "source_sheet": "data",
                    "excel_row": excel_row, "column_index": col[cname], "column_name": cname,
                    "raw_value": rawv, "hub_project_id": r[col["Id"]],
                })

        dur = (ey - sy + 1) if (sy is not None and ey is not None and ey >= sy) else None

        def g(c):
            v = r[col[c]]
            return "" if v is None else str(v)

        out_rows.append({
            "hub_project_id": r[col["Id"]], "source_row_index": excel_row,
            "funder_name": g("Funder Name"), "funder_country_raw": g("Funder Country"),
            "funder_iso3": f_iso3 or "", "funder_country_status": f_status,
            "institution_name": g("Institution Name"),
            "institution_country_raw": g("Institution Country"),
            "institution_iso3": i_iso3 or "", "institution_country_status": i_status,
            "start_year": "" if sy is None else sy, "start_year_status": sy_st,
            "end_year": "" if ey is None else ey, "end_year_status": ey_st,
            "duration_years": "" if dur is None else dur,
            "amount_usd": "" if usd is None else usd, "amount_usd_status": usd_st,
            "amount_usd_raw": usd_raw,
            "amount_eur": "" if eur is None else eur, "amount_eur_status": eur_st,
            "amount_eur_raw": eur_raw,
            "sector": g("Sector"), "sector_subcategory": g("Sector Subcategory"),
            "research_area": g("Research Area"), "research_subcategory": g("Research Subcategory"),
            "disease": g("Disease"), "infectious_agent": g("Infectious Agent"),
            "individual_infectious_agent": g("Individual Infectious Agent"),
            "product_name": g("Product Name"), "project_group": g("Project Group"),
            "acronym": g("Acronym"),
            "source_file": os.path.basename(SOURCE_XLSX), "source_sheet": "data",
            "header_row": HEADER_ROW, "retrieval_date": RETRIEVAL_DATE,
        })

    with open(os.path.join(OUT, "hub_projects_harmonised.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=proj_fields)
        w.writeheader()
        w.writerows(out_rows)
    with open(os.path.join(OUT, "hub_unparsed_values.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["source_file", "source_sheet", "excel_row", "column_index",
                                          "column_name", "raw_value", "hub_project_id"])
        w.writeheader()
        w.writerows(unparsed_log)
    assert len(out_rows) == n, f"row count changed: {len(out_rows)} != {n}"

    # country x year totals (funder and host attribution)
    # Annualised = amount spread evenly over start..end years; start-year attribution
    # is written alongside.
    agg = collections.defaultdict(lambda: {"n_projects": 0, "usd_startyear": 0.0, "usd_annualised": 0.0})
    for row in out_rows:
        if row["amount_usd_status"] != "ok":
            continue
        usd = float(row["amount_usd"])
        sy, ey = row["start_year"], row["end_year"]
        for basis, iso3 in (("funder", row["funder_iso3"]), ("institution", row["institution_iso3"])):
            if not iso3:
                continue
            if sy != "":
                k = (basis, iso3, int(sy))
                agg[k]["n_projects"] += 1
                agg[k]["usd_startyear"] += usd
            if sy != "" and ey != "" and int(ey) >= int(sy):
                years = list(range(int(sy), int(ey) + 1))
                for y in years:
                    agg[(basis, iso3, y)]["usd_annualised"] += usd / len(years)
    with open(os.path.join(OUT, "hub_investment_by_country_year.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["attribution_basis", "iso3", "year", "n_projects_starting",
                    "usd_nominal_startyear", "usd_nominal_annualised", "source_file", "retrieval_date"])
        for (basis, iso3, y) in sorted(agg):
            v = agg[(basis, iso3, y)]
            w.writerow([basis, iso3, y, v["n_projects"], round(v["usd_startyear"], 2),
                        round(v["usd_annualised"], 2), os.path.basename(SOURCE_XLSX), RETRIEVAL_DATE])

    # Country totals; the merge script selects the study countries.
    hub_funder = {r["funder_iso3"] for r in out_rows if r["funder_iso3"]}
    hub_inst = {r["institution_iso3"] for r in out_rows if r["institution_iso3"]}
    fu, fu_usd = collections.Counter(), collections.Counter()
    iu, iu_usd = collections.Counter(), collections.Counter()
    fu_years, iu_years = collections.defaultdict(set), collections.defaultdict(set)
    for r in out_rows:
        usd = float(r["amount_usd"]) if r["amount_usd_status"] == "ok" else 0.0
        if r["funder_iso3"]:
            fu[r["funder_iso3"]] += 1
            fu_usd[r["funder_iso3"]] += usd
            if r["start_year"] != "":
                fu_years[r["funder_iso3"]].add(int(r["start_year"]))
        if r["institution_iso3"]:
            iu[r["institution_iso3"]] += 1
            iu_usd[r["institution_iso3"]] += usd
            if r["start_year"] != "":
                iu_years[r["institution_iso3"]].add(int(r["start_year"]))

    analysis_set = sorted(hub_funder | hub_inst)
    cov_rows = []
    for iso3 in analysis_set:
        fy, iy = sorted(fu_years[iso3]), sorted(iu_years[iso3])
        cov_rows.append({
            "iso3": iso3,
            "in_hub_as_funder": iso3 in hub_funder, "n_projects_funded": fu[iso3],
            "usd_nominal_funded": round(fu_usd[iso3], 2),
            "funder_year_min": fy[0] if fy else "", "funder_year_max": fy[-1] if fy else "",
            "funder_n_distinct_start_years": len(fy),
            "in_hub_as_institution": iso3 in hub_inst, "n_projects_hosted": iu[iso3],
            "usd_nominal_hosted": round(iu_usd[iso3], 2),
            "institution_year_min": iy[0] if iy else "", "institution_year_max": iy[-1] if iy else "",
            "institution_n_distinct_start_years": len(iy),
        })
    with open(os.path.join(OUT, "hub_by_country.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(cov_rows[0].keys()))
        w.writeheader()
        w.writerows(cov_rows)

    total_usd = sum(float(r["amount_usd"]) for r in out_rows if r["amount_usd_status"] == "ok")
    prov = {
        "run_started": run_started,
        "source_urls": SOURCE_URLS,
        "source_file": {
            "basename": os.path.basename(SOURCE_XLSX),
            "bytes": os.path.getsize(SOURCE_XLSX),
            "sha256": sha256(SOURCE_XLSX),
            "export_metadata_sheet": {k: (str(v) if v is not None else None) for k, v in meta.items()},
            "retrieval_date": RETRIEVAL_DATE,
        },
        "unit_of_observation": "one funded project / grant (column 'Id')",
        "n_projects": n,
        "total_amount_usd_nominal": round(total_usd, 2),
        "currency_note": "Nominal award values as supplied by the Hub; not inflation- or PPP-adjusted.",
        "n_countries": len(analysis_set),
        "coverage": {
            "n_present_as_funder": sum(1 for r in cov_rows if r["in_hub_as_funder"]),
            "n_present_as_institution": sum(1 for r in cov_rows if r["in_hub_as_institution"]),
        },
        "checks": {"unmapped_country_names": len(unmapped), "unparsed_values": len(unparsed_log),
                   "rows_in": n, "rows_out": len(out_rows)},
    }
    with open(os.path.join(OUT, "HUB_PROVENANCE.json"), "w", encoding="utf-8") as f:
        json.dump(prov, f, indent=2, ensure_ascii=False)

    print(f"projects={n}  total USD={total_usd:,.0f}  unparsed={len(unparsed_log)}")
    print(f"countries n={len(analysis_set)}; hosted-investment records for "
          f"{prov['coverage']['n_present_as_institution']}")


def main():
    os.chdir(ROOT)
    clean_projects()


if __name__ == "__main__":
    main()
